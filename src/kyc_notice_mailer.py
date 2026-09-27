"""Bank-wise KYC notice mailer.

Upload a pending KYC sheet and a bank e-mail directory, review the matched
recipients, then send one mail per bank straight through the NIC (mgovcloud)
mail server over SMTP - no browser extension involved.

Each mail body is the bank's Word notice rendered as HTML, stopping before the
officer's stamp, signature and name. Only the PDF notice and the KYC Excel are
attached; Word files are never attached.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import json
import secrets
import threading
from datetime import datetime
from difflib import SequenceMatcher
from email.message import EmailMessage
from email.utils import formatdate, make_msgid
from html import escape
import io
import os
from pathlib import Path
import re
import smtplib
import subprocess
import sys
import time
import zipfile
from typing import Dict, List, Optional, Sequence, Tuple

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from docx.oxml.ns import qn

from src.bankwise_kyc_notices import build_bank_bundles, build_payer_remitter_bundles, pending_rows
from src.kyc_prabodh import build_prabodh_bundles
from src.mail_screenshot import (
    build_sent_mail_html,
    inline_cid_images,
    records_folder,
    render_html_files_to_png,
    screenshot_filename,
)
from src.gujarat_account_formatter import _current_notice_date

SMTP_HOST = os.environ.get("NIC_SMTP_HOST", "smtp.mgovcloud.in")
SMTP_PORT = int(os.environ.get("NIC_SMTP_PORT", "465"))
DEFAULT_SUBJECT = "Provision of Know Your Customer (KYC) Details of Account Holder(s)"
DEFAULT_PR_SUBJECT = "Request for Providing Last One Month Bank Account Statement (Payer / Remitter)"

# Notice type keys
NOTICE_TYPE_KYC = "KYC Notice"
NOTICE_TYPE_PR = "Payer / Remitter Statement Request"
NOTICE_TYPE_PRABODH = "KYC Prabodh Notice"
DEFAULT_EMAIL_SHEET = Path(__file__).resolve().parent.parent / "BANK MAIL.xlsx"
SEND_LOG_PATH = Path(__file__).resolve().parent.parent / "data" / "kyc_mail_sent_log.csv"
ACCESS_FILE = Path(__file__).resolve().parent.parent / "data" / "kyc_mailer_access.json"
SEND_LOG_FIELDS = ["sent_at", "bank_name", "subject", "recipients", "attachments", "status", "error", "sender", "sender_name"]
DEFAULT_SENDER_NAME = "HELPLINE UNIT-16 CYBER CRIME CELL (GoG Home Dept)"
ATTACHMENT_TYPES = {
    ".pdf": ("application", "pdf"),
    ".xlsx": ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_EMU_PER_PX = 9525
_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"


# --------------------------------------------------------------------------- #
# Word notice -> e-mail HTML
# --------------------------------------------------------------------------- #
def _element_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _element_images(element) -> List[Tuple[str, int, int]]:
    images = []
    for drawing in element.iter(qn("w:drawing")):
        blip = next(drawing.iter(_BLIP), None)
        extent = next(drawing.iter(_EXTENT), None)
        if blip is None:
            continue
        width = int(extent.get("cx")) // _EMU_PER_PX if extent is not None else 100
        height = int(extent.get("cy")) // _EMU_PER_PX if extent is not None else 100
        images.append((blip.get(qn("r:embed")), width, height))
    return images


def _run_html(run) -> str:
    parts = []
    for child in run.iterchildren():
        if child.tag == qn("w:t"):
            text = escape(child.text or "")
            parts.append(re.sub(r"(?<= ) |^ ", "&nbsp;", text))
        elif child.tag == qn("w:tab"):
            parts.append("&nbsp;" * 8)
        elif child.tag in (qn("w:br"), qn("w:cr")):
            parts.append("<br>")
    html = "".join(parts)
    if not html:
        return ""

    props = run.find(qn("w:rPr"))
    styles = []
    if props is not None:
        if props.find(qn("w:b")) is not None and props.find(qn("w:b")).get(qn("w:val")) not in ("0", "false"):
            html = f"<b>{html}</b>"
        if props.find(qn("w:i")) is not None and props.find(qn("w:i")).get(qn("w:val")) not in ("0", "false"):
            html = f"<i>{html}</i>"
        underline = props.find(qn("w:u"))
        if underline is not None and underline.get(qn("w:val")) not in (None, "none"):
            html = f"<u>{html}</u>"
        color = props.find(qn("w:color"))
        if color is not None and re.fullmatch(r"[0-9A-Fa-f]{6}", color.get(qn("w:val")) or ""):
            styles.append(f"color:#{color.get(qn('w:val'))}")
        size = props.find(qn("w:sz"))
        if size is not None and (size.get(qn("w:val")) or "").isdigit():
            styles.append(f"font-size:{int(size.get(qn('w:val'))) / 2:g}pt")
    return f'<span style="{";".join(styles)}">{html}</span>' if styles else html


def _paragraph_runs_html(paragraph) -> str:
    return "".join(_run_html(run) for run in paragraph.iter(qn("w:r")))


def _paragraph_align(paragraph) -> str:
    jc = paragraph.find(f"{qn('w:pPr')}/{qn('w:jc')}")
    value = jc.get(qn("w:val")) if jc is not None else "left"
    return {"both": "justify", "center": "center", "right": "right", "end": "right"}.get(value, "left")


def _paragraph_html(paragraph, image_cids: Dict[str, str], continuation: str = "") -> str:
    align = _paragraph_align(paragraph)
    images = _element_images(paragraph)
    body = _paragraph_runs_html(paragraph) + continuation
    style = f"margin:0 0 8px 0;text-align:{align}"

    if images:
        rel_id, width, height = images[0]
        image = (
            f'<img src="cid:{image_cids[rel_id]}" width="{width}" height="{height}" '
            f'alt="" style="display:block">'
        )
        if not _element_text(paragraph).strip():
            return f'<p style="{style}">{image}</p>'
        # Letterhead: logo on the left, office address centred beside it.
        return (
            '<table role="presentation" style="border-collapse:collapse;width:100%;margin:0">'
            f'<tr><td style="width:{width + 12}px;vertical-align:middle;padding:0">{image}</td>'
            f'<td style="text-align:center;vertical-align:middle;padding:0">{body}</td></tr></table>'
        )

    # A mid-line tab (reference number ... date) becomes a left/right split.
    text = _element_text(paragraph)
    tab_runs = [run for run in paragraph.iter(qn("w:r")) if run.find(qn("w:tab")) is not None]
    if tab_runs and text.strip() and not text.startswith("\t") and not _paragraph_starts_with_tab(paragraph):
        left, right = [], []
        target = left
        for run in paragraph.iter(qn("w:r")):
            if run.find(qn("w:tab")) is not None and not _run_has_text(run):
                target = right
                continue
            target.append(_run_html(run))
        return (
            '<table role="presentation" style="border-collapse:collapse;width:100%;margin:0 0 8px 0">'
            f'<tr><td style="padding:0;text-align:left">{"".join(left)}</td>'
            f'<td style="padding:0;text-align:right">{"".join(right)}</td></tr></table>'
        )

    return f'<p style="{style}">{body or "&nbsp;"}</p>'


def _run_has_text(run) -> bool:
    return any((node.text or "").strip() for node in run.iter(qn("w:t")))


def _paragraph_starts_with_tab(paragraph) -> bool:
    for run in paragraph.iter(qn("w:r")):
        for child in run.iterchildren():
            if child.tag == qn("w:tab"):
                return True
            if child.tag == qn("w:t") and (child.text or ""):
                return False
    return False


def _table_html(table) -> str:
    rows = []
    for row_index, row in enumerate(table.iter(qn("w:tr"))):
        cells = []
        for cell in row.iter(qn("w:tc")):
            paragraphs = cell.findall(qn("w:p"))
            align = _paragraph_align(paragraphs[0]) if paragraphs else "left"
            content = "<br>".join(_paragraph_runs_html(p) for p in paragraphs)
            tag = "th" if row_index == 0 else "td"
            background = "background:#D9D9D9;" if row_index == 0 else ""
            cells.append(
                f'<{tag} style="{background}border:1px solid #9E9E9E;padding:5px 8px;'
                f'text-align:{"center" if row_index == 0 else align};vertical-align:middle">'
                f"{content}</{tag}>"
            )
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return (
        '<table style="border-collapse:collapse;margin:4px 0 12px 0;'
        'font-family:Calibri,Arial,sans-serif;font-size:10pt">' + "".join(rows) + "</table>"
    )


def notice_docx_to_email(docx_bytes: bytes) -> Tuple[str, Dict[str, Tuple[str, bytes]]]:
    """Render the Word notice as e-mail HTML, excluding stamp, signature and officer.

    Returns the HTML and ``{content_id: (image_subtype, bytes)}`` for inline images.
    """
    document = Document(io.BytesIO(docx_bytes))
    body = document.element.body
    elements = [el for el in body.iterchildren() if el.tag in (qn("w:p"), qn("w:tbl"))]

    image_cids: Dict[str, str] = {}
    inline_images: Dict[str, Tuple[str, bytes]] = {}
    html_parts = []
    seen_text = False
    absorbed = set()
    for index, element in enumerate(elements):
        text = _element_text(element)
        images = _element_images(element) if element.tag == qn("w:p") else []
        # The stamp/signature image is the first image-only paragraph after the
        # notice text starts; the officer's name block follows it.
        if seen_text and images and not text.strip():
            break
        if "Keshvala" in text or text.strip().startswith("Superintendent of Police"):
            break

        for rel_id, _, _ in images:
            if rel_id not in image_cids:
                part = document.part.related_parts[rel_id]
                cid = make_msgid(domain="kyc.notice")[1:-1]
                image_cids[rel_id] = cid
                subtype = part.content_type.split("/")[-1]
                inline_images[cid] = (subtype, part.blob)

        if index in absorbed:
            continue
        if element.tag == qn("w:tbl"):
            html_parts.append(_table_html(element))
        elif images and text.strip():
            # Letterhead: the address lines after the logo paragraph sit beside the logo.
            continuation = []
            for follower_index in range(index + 1, len(elements)):
                follower = elements[follower_index]
                if follower.tag != qn("w:p") or not _element_text(follower).strip():
                    break
                continuation.append("<br>" + _paragraph_runs_html(follower))
                absorbed.add(follower_index)
            html_parts.append(_paragraph_html(element, image_cids, "".join(continuation)))
        else:
            html_parts.append(_paragraph_html(element, image_cids))
        seen_text = seen_text or bool(text.strip())

    # Drop trailing empty paragraphs.
    while html_parts and html_parts[-1].endswith(">&nbsp;</p>"):
        html_parts.pop()

    html = (
        '<div style="font-family:\'Times New Roman\',Times,serif;font-size:12pt;'
        'color:#000;max-width:720px">' + "".join(html_parts) + "</div>"
    )
    return html, inline_images


def notice_subject_from_docx(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    for paragraph in document.paragraphs:
        match = re.search(r"Subject\s*:\s*(.+)", paragraph.text)
        if match:
            return match.group(1).strip().rstrip(".")
    return DEFAULT_SUBJECT


def build_subject(base: str, period: str, bank_name: str) -> str:
    base = base.strip().rstrip(".")
    period = period.strip()
    suffix = f"{period}- {bank_name}" if period else bank_name
    return f"{base} ({suffix})"


# --------------------------------------------------------------------------- #
# Bank e-mail directory
# --------------------------------------------------------------------------- #
def extract_emails(value) -> List[str]:
    unique, seen = [], set()
    for address in _EMAIL_RE.findall(str(value or "")):
        address = address.strip(".")
        if address.casefold() not in seen:
            seen.add(address.casefold())
            unique.append(address)
    return unique


def normalize_bank_name(value) -> str:
    value = str(value or "").lower()
    value = value.replace("’", "'").replace("people's", "peoples")
    value = re.sub(r"\(\s*including\b[^)]*\)", " ", value)
    value = value.replace("&", " and ")
    value = re.sub(r"\bco[\s.-]*op(?:erative)?\b", " cooperative ", value)
    value = re.sub(r"\bcoop(?:erative)?\b", " cooperative ", value)
    value = re.sub(r"\bdist\b", "district", value)
    value = re.sub(r"[^\w\s]", " ", value)
    stop = {"the", "and", "of", "ltd", "limited", "co", "company", "including", "upi"}
    words = []
    for word in value.split():
        if word not in stop and word not in words:
            words.append(word)
    return " ".join(words)


# KYC-sheet bank name -> the exact row (CATEGORY) to use in the e-mail sheet.
EXACT_SHEET_ROWS = {
    "AU Bank": "AU SMALL FINANCE BANK",
    "THE GUJARAT STATE COOPERATIVE BANK": "THE GUJARAT STATE COOPERATIVE BANK",
}


def _prefer_non_upi(rows: List[Dict]) -> List[Dict]:
    """UPI rows are used only when the bank has no other row in the sheet."""
    regular = [row for row in rows if "UPI" not in re.findall(r"[A-Za-z]+", row["bank_name"].upper())]
    return regular or rows


def _sheet_name_key(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


_GENERIC_TOKENS = {
    "bank", "small", "finance", "commercial", "cooperative", "gramin",
    "sahakari", "peoples", "urban", "nagrik", "nagarik", "district", "central",
}


_BANK_TYPES = {
    "gramin": "gramin",
    "cooperative": "cooperative",
    "sahakari": "cooperative",
    "nagrik": "cooperative",
    "nagarik": "cooperative",
    "small": "small finance",
    "finance": "small finance",
}


def _bank_types(tokens: set) -> set:
    return {_BANK_TYPES[token] for token in tokens if token in _BANK_TYPES}


def detect_directory_columns(columns: List[str]) -> Tuple[Optional[str], Optional[str]]:
    normalized = {re.sub(r"[^a-z]", "", str(c).lower()): c for c in columns}
    bank_col = next(
        (normalized[k] for k in ("bankname", "category", "bank", "name", "nameofbank") if k in normalized),
        None,
    )
    email_col = next(
        (normalized[k] for k in ("mailids", "emailids", "email", "emails", "mailid", "emailid", "mail") if k in normalized),
        None,
    )
    return bank_col, email_col


def load_email_directory(df: pd.DataFrame, bank_col: str, email_col: str) -> List[Dict]:
    rows = []
    for bank, emails in zip(df[bank_col], df[email_col]):
        addresses = extract_emails(emails)
        if str(bank or "").strip() and addresses:
            rows.append({"bank_name": str(bank).strip(), "key": normalize_bank_name(bank), "emails": addresses})
    return rows


def _has_similar_token(token: str, tokens: set) -> bool:
    return token in tokens or any(
        len(token) > 3 and SequenceMatcher(None, token, other).ratio() >= 0.85 for other in tokens
    )


def match_bank_emails(bank_name: str, directory: List[Dict]) -> Dict:
    """Find the directory entry for a bank. Exact normalized names are merged;
    otherwise the closest name that shares every distinctive word is used."""
    fixed = {normalize_bank_name(k): _sheet_name_key(v) for k, v in EXACT_SHEET_ROWS.items()}
    key = normalize_bank_name(bank_name)
    if key in fixed:
        rows = [row for row in directory if _sheet_name_key(row["bank_name"]) == fixed[key]]
        if rows:
            emails = extract_emails(";".join(";".join(row["emails"]) for row in rows))
            return {"matched_name": rows[0]["bank_name"], "emails": emails, "match": "Exact"}

    exact = _prefer_non_upi([row for row in directory if row["key"] == key])
    if exact:
        emails = extract_emails(";".join(";".join(row["emails"]) for row in exact))
        return {"matched_name": exact[0]["bank_name"], "emails": emails, "match": "Exact"}

    query_tokens = set(key.split())
    identity = query_tokens - _GENERIC_TOKENS
    best, best_score = None, 0.0
    for row in directory:
        candidate_tokens = set(row["key"].split())
        if identity and not all(_has_similar_token(token, candidate_tokens) for token in identity):
            continue
        extra = [t for t in candidate_tokens - _GENERIC_TOKENS if not _has_similar_token(t, query_tokens)]
        if len(extra) > 1:
            continue
        # A Gramin bank must not match a Cooperative bank of the same state, etc.
        query_types, candidate_types = _bank_types(query_tokens), _bank_types(candidate_tokens)
        if query_types and candidate_types and query_types != candidate_types:
            continue
        token_score = len(query_tokens & candidate_tokens) / max(len(query_tokens | candidate_tokens), 1)
        score = max(SequenceMatcher(None, key, row["key"]).ratio(), token_score)
        if identity and not extra:
            score = max(score, 0.9)
        if score > best_score:
            best, best_score = row, score
    if best is not None and best_score >= 0.6:
        rows = _prefer_non_upi([row for row in directory if row["key"] == best["key"]])
        return {
            "matched_name": rows[0]["bank_name"],
            "emails": extract_emails(";".join(";".join(row["emails"]) for row in rows)),
            "match": f"Closest ({best_score:.0%})",
        }
    return {"matched_name": "", "emails": [], "match": "Not found"}


# --------------------------------------------------------------------------- #
# Message building and sending
# --------------------------------------------------------------------------- #
def build_message(
    sender: str,
    recipients: List[str],
    subject: str,
    html: str,
    inline_images: Dict[str, Tuple[str, bytes]],
    attachments: Dict[str, bytes],
    cc: Optional[List[str]] = None,
) -> EmailMessage:
    for filename in attachments:
        if Path(filename).suffix.lower() not in ATTACHMENT_TYPES:
            raise ValueError(f"Refusing to attach '{filename}': only PDF and Excel files are sent.")

    message = EmailMessage()
    message["From"] = sender
    message["To"] = ", ".join(recipients)
    if cc:
        message["Cc"] = ", ".join(cc)
    message["Subject"] = subject
    message["Date"] = formatdate(localtime=True)
    message["Message-ID"] = make_msgid(domain=sender.split("@")[-1])

    plain = re.sub(r"<br\s*/?>|</p>|</tr>|</table>", "\n", html)
    plain = re.sub(r"</t[dh]>", "\t", plain)
    plain = re.sub(r"<[^>]+>", "", plain).replace("&nbsp;", " ")
    plain = re.sub(r"\n{3,}", "\n\n", re.sub(r"&amp;", "&", plain)).strip()
    message.set_content(plain)
    message.add_alternative(f"<html><body>{html}</body></html>", subtype="html")
    html_part = message.get_payload()[1]
    for cid, (subtype, data) in inline_images.items():
        html_part.add_related(data, maintype="image", subtype=subtype, cid=f"<{cid}>", disposition="inline")

    for filename, data in attachments.items():
        maintype, subtype = ATTACHMENT_TYPES[Path(filename).suffix.lower()]
        message.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    return message


def open_smtp(sender: str, password: str, host: str = SMTP_HOST, port: int = SMTP_PORT) -> smtplib.SMTP:
    if port == 465:
        connection = smtplib.SMTP_SSL(host, port, timeout=60)
    else:
        connection = smtplib.SMTP(host, port, timeout=60)
        connection.ehlo()
        if connection.has_extn("starttls"):
            connection.starttls()
            connection.ehlo()
        elif host not in ("localhost", "127.0.0.1"):
            connection.close()
            raise ConnectionError(f"{host}:{port} does not offer encryption; refusing to send the password.")
    connection.login(sender, password)
    return connection


def mail_attachments(bundle: Dict) -> Dict[str, bytes]:
    return {
        name: data
        for name, data in bundle["files"].items()
        if Path(name).suffix.lower() in ATTACHMENT_TYPES
    }


def read_send_log() -> pd.DataFrame:
    if not SEND_LOG_PATH.exists():
        return pd.DataFrame(columns=SEND_LOG_FIELDS)
    return pd.read_csv(SEND_LOG_PATH, dtype=str, keep_default_na=False)


def _upgrade_send_log_header() -> None:
    """Keep an older log readable after new columns were added."""
    if not SEND_LOG_PATH.exists():
        return
    with SEND_LOG_PATH.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
        header = rows and set(rows[0]) or set()
    if rows and not set(SEND_LOG_FIELDS) <= header:
        with SEND_LOG_PATH.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=SEND_LOG_FIELDS)
            writer.writeheader()
            for row in rows:
                writer.writerow({field: row.get(field, "") or "" for field in SEND_LOG_FIELDS})


def append_send_log(entry: Dict) -> None:
    SEND_LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    _upgrade_send_log_header()
    new_file = not SEND_LOG_PATH.exists()
    with SEND_LOG_PATH.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=SEND_LOG_FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow({field: entry.get(field, "") for field in SEND_LOG_FIELDS})


def save_sent_mail_screenshots(
    entries: List[Dict],
    bundle_by_bank: Dict[str, Dict],
    subject_base: str,
    period: str,
    notice_date,
) -> List[Path]:
    """Render one mailbox-style PNG per sent notice into the records folder."""
    folder = records_folder(notice_date)
    if any(" KYC Prabodh Notice.docx" in filename
           for bundle in bundle_by_bank.values() for filename in bundle.get("files", {})):
        folder = folder / "KYC Prabodh"
    pages: List[Tuple[str, Path]] = []
    for entry in entries:
        bundle = bundle_by_bank.get(entry["bank_name"])
        if bundle is None:
            continue
        docx_bytes = next(d for n, d in bundle["files"].items() if n.endswith(".docx"))
        html, images = _cached_email(docx_bytes)
        attachments = [(name, len(data)) for name, data in mail_attachments(bundle).items()]
        page = build_sent_mail_html(
            entry.get("subject") or build_subject(subject_base, period, entry["bank_name"]),
            entry.get("sender_name") or DEFAULT_SENDER_NAME,
            entry.get("sender", ""),
            extract_emails(entry.get("recipients", "")),
            entry.get("sent_at", ""),
            inline_cid_images(html, images),
            attachments,
        )
        pages.append((page, folder / screenshot_filename(entry["bank_name"])))
    return render_html_files_to_png(pages)


def screenshots_zip_bytes(paths: Sequence[Path]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in paths:
            archive.write(path, arcname=path.name)
    return buffer.getvalue()


def whatsapp_profile_folder(officer: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", officer or "").strip("_").lower()
    return f"whatsapp_profile_{slug}" if slug else "whatsapp_profile"


def run_whatsapp(arguments: List[str], on_event, officer: str = "") -> int:
    """Run the WhatsApp helper as its own process and stream its JSON progress."""
    environment = {**os.environ, "WHATSAPP_PROFILE": whatsapp_profile_folder(officer)}
    process = subprocess.Popen(
        [sys.executable, str(Path(__file__).resolve().parent / "whatsapp_sender.py"), *arguments],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        env=environment,
    )
    for line in process.stdout:
        line = line.strip()
        if not line:
            continue
        try:
            on_event(json.loads(line))
        except ValueError:
            on_event({"status": "log", "message": line})
    return process.wait()


def already_sent_subjects() -> set:
    log = read_send_log()
    return set(log.loc[log["status"] == "sent", "subject"]) if not log.empty else set()


# --------------------------------------------------------------------------- #
# Streamlit page
# --------------------------------------------------------------------------- #
def _login_failed_message(exc: smtplib.SMTPAuthenticationError) -> str:
    detail = exc.smtp_error.decode(errors="replace") if isinstance(exc.smtp_error, bytes) else str(exc.smtp_error)
    return (
        f"Login refused by {SMTP_HOST} ({exc.smtp_code} {detail.strip()}). "
        "NIC/mgovcloud accounts with two-factor login do not accept the normal web password here: "
        "open mgovcloud Accounts -> Security -> App Passwords, generate one, and paste it in the password box. "
        "If that is also refused, ask the mail admin to enable SMTP access for this mailbox."
    )


def _secret(name: str) -> str:
    try:
        return str(st.secrets.get("nic_mail", {}).get(name, ""))
    except Exception:
        return ""


@st.cache_data(show_spinner=False)
def _cached_bundles(file_bytes: bytes, notice_date) -> List[Dict]:
    df = pd.read_excel(io.BytesIO(file_bytes), dtype=str, keep_default_na=False)
    return build_bank_bundles(pending_rows(df), notice_date)


@st.cache_data(show_spinner=False)
def _cached_pr_bundles(file_bytes: bytes, notice_date) -> List[Dict]:
    df = pd.read_excel(io.BytesIO(file_bytes), dtype=str, keep_default_na=False)
    return build_payer_remitter_bundles(pending_rows(df), notice_date)


@st.cache_data(show_spinner=False)
def _cached_prabodh_bundles(file_bytes: bytes, notice_date) -> List[Dict]:
    df = pd.read_excel(io.BytesIO(file_bytes), dtype=str, keep_default_na=False)
    return build_prabodh_bundles(df, notice_date)


@st.cache_data(show_spinner=False)
def _cached_email(docx_bytes: bytes):
    return notice_docx_to_email(docx_bytes)


# --------------------------------------------------------------------------- #
# Page password
# --------------------------------------------------------------------------- #
_PBKDF2_ITERATIONS = 240_000
_MAX_FAILED_ATTEMPTS = 5
_LOCKOUT_SECONDS = 300
_failed_attempts = {"count": 0, "locked_until": 0.0}
_failed_lock = threading.Lock()


def hash_page_password(password: str, salt: Optional[bytes] = None) -> Dict:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, _PBKDF2_ITERATIONS)
    return {"salt": salt.hex(), "hash": digest.hex(), "iterations": _PBKDF2_ITERATIONS}


def verify_page_password(password: str, record: Dict) -> bool:
    digest = hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), bytes.fromhex(record["salt"]), int(record["iterations"])
    )
    return hmac.compare_digest(digest.hex(), record["hash"])


def load_access_record(path: Optional[Path] = None) -> Optional[Dict]:
    path = path or ACCESS_FILE
    try:
        record = json.loads(path.read_text(encoding="utf-8"))
        return record if {"salt", "hash", "iterations"} <= set(record) else None
    except (OSError, ValueError):
        return None


def save_access_record(password: str, path: Optional[Path] = None) -> None:
    path = path or ACCESS_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(hash_page_password(password)), encoding="utf-8")


def _password_rule_error(password: str, confirm: str) -> Optional[str]:
    if len(password) < 8:
        return "Password must be at least 8 characters."
    if password != confirm:
        return "The two passwords do not match."
    return None


def _require_page_password() -> bool:
    """Show the lock screen until this browser session enters the page password."""
    record = load_access_record()
    if record and st.session_state.get("kyc_mailer_unlocked") == record["hash"]:
        return True

    st.title("KYC Notice Mailer (NIC Mail)")
    if record is None:
        st.info("This page is protected. Create the password that will be needed to open it.")
        with st.form("kyc_mailer_create_password"):
            new_password = st.text_input("New page password", type="password")
            confirm = st.text_input("Confirm password", type="password")
            if st.form_submit_button("Set password", type="primary"):
                error = _password_rule_error(new_password, confirm)
                if error:
                    st.error(error)
                else:
                    save_access_record(new_password)
                    st.session_state["kyc_mailer_unlocked"] = load_access_record()["hash"]
                    st.rerun()
        return False

    st.warning("This page is locked. Enter the page password to continue.")
    with st.form("kyc_mailer_unlock"):
        password = st.text_input("Page password", type="password")
        submitted = st.form_submit_button("Unlock", type="primary")
    if submitted:
        with _failed_lock:
            wait = int(_failed_attempts["locked_until"] - time.time())
            if wait > 0:
                st.error(f"Too many wrong attempts. Try again in {wait // 60 + 1} minute(s).")
                return False
            if verify_page_password(password, record):
                _failed_attempts.update(count=0, locked_until=0.0)
                st.session_state["kyc_mailer_unlocked"] = record["hash"]
                st.rerun()
            _failed_attempts["count"] += 1
            if _failed_attempts["count"] >= _MAX_FAILED_ATTEMPTS:
                _failed_attempts.update(count=0, locked_until=time.time() + _LOCKOUT_SECONDS)
                st.error("Too many wrong attempts. The page is blocked for 5 minutes.")
            else:
                st.error(f"Wrong password. {_MAX_FAILED_ATTEMPTS - _failed_attempts['count']} attempt(s) left.")
    return False


def _render_password_controls() -> None:
    col1, col2 = st.columns([1, 5])
    with col1:
        if st.button("Lock page", key="kyc_mailer_lock"):
            st.session_state.pop("kyc_mailer_unlocked", None)
            st.rerun()
    with col2:
        with st.expander("Change page password"):
            with st.form("kyc_mailer_change_password", clear_on_submit=True):
                current = st.text_input("Current password", type="password")
                new_password = st.text_input("New password", type="password")
                confirm = st.text_input("Confirm new password", type="password")
                if st.form_submit_button("Change password"):
                    record = load_access_record()
                    error = _password_rule_error(new_password, confirm)
                    if not record or not verify_page_password(current, record):
                        st.error("Current password is wrong.")
                    elif error:
                        st.error(error)
                    else:
                        save_access_record(new_password)
                        st.session_state["kyc_mailer_unlocked"] = load_access_record()["hash"]
                        st.success("Page password changed.")


def _entries_by_send_day(entries: List[Dict], fallback) -> List[Tuple[object, List[Dict]]]:
    grouped: Dict[object, List[Dict]] = {}
    for entry in entries:
        try:
            day = datetime.strptime(entry["sent_at"], "%Y-%m-%d %H:%M:%S").date()
        except (KeyError, ValueError):
            day = fallback
        grouped.setdefault(day, []).append(entry)
    return sorted(grouped.items(), key=lambda item: str(item[0]))


def _render_records_section(bundle_by_bank, subject_base, period, notice_date, sender_name, sender, kyc_bytes, is_payer_remitter=False, is_prabodh=False) -> None:
    st.subheader("Sent-mail records & WhatsApp")
    folder = records_folder(notice_date)
    if is_prabodh:
        folder = folder / "KYC Prabodh"
    log = read_send_log()
    sent_rows = (
        log[(log["status"] == "sent") & log["bank_name"].isin(bundle_by_bank)
            & log["subject"].isin({build_subject(subject_base, period, bank) for bank in bundle_by_bank})].drop_duplicates("bank_name", keep="last")
        if not log.empty
        else pd.DataFrame(columns=SEND_LOG_FIELDS)
    )
    existing = sorted(folder.glob("*.png")) if folder.exists() else []
    st.caption(f"{len(sent_rows)} mail(s) sent for this file - {len(existing)} screenshot(s) saved in {folder}")

    col1, col2 = st.columns([1, 2])
    with col1:
        rebuild = st.button("Create / refresh screenshots", key="kyc_mailer_shots", disabled=sent_rows.empty)
    if rebuild:
        entries = sent_rows.to_dict("records")
        for entry in entries:
            # Older log rows were written before the sender columns existed.
            entry["sender_name"] = entry.get("sender_name") or sender_name
            entry["sender"] = entry.get("sender") or sender
        with st.spinner(f"Rendering {len(entries)} screenshot(s)..."):
            try:
                for send_day, day_entries in _entries_by_send_day(entries, notice_date):
                    # Rebuild each notice with the date it was actually sent on,
                    # so the screenshot matches the mail the bank received.
                    day_bundles = (
                        bundle_by_bank
                        if send_day == notice_date
                        else {
                            bundle["bank_name"]: bundle
                            for bundle in (
                                _cached_prabodh_bundles(kyc_bytes, send_day)
                                if is_prabodh
                                else _cached_pr_bundles(kyc_bytes, send_day) if is_payer_remitter
                                else _cached_bundles(kyc_bytes, send_day)
                            )
                            if bundle.get("files")
                        }
                    )
                    save_sent_mail_screenshots(day_entries, day_bundles, subject_base, period, notice_date)
                existing = sorted(folder.glob("*.png"))
                st.success(f"{len(existing)} screenshot(s) saved in {folder}")
            except Exception as exc:
                st.error(f"Could not render the screenshots: {exc}")
        existing = sorted(folder.glob("*.png"))

    if not existing:
        st.info("Send the notices (or press the button above) to build the record screenshots.")
        return

    with col2:
        st.download_button(
            f"Download all {len(existing)} screenshots (ZIP)",
            screenshots_zip_bytes(existing),
            file_name=f"KYC sent mails {notice_date}.zip",
            key="kyc_mailer_zip",
        )
    preview = st.selectbox("Preview screenshot", [path.name for path in existing], key="kyc_mailer_shot_preview")
    st.image(str(folder / preview), width="stretch")

    st.markdown("**Forward on WhatsApp** - one message per bank, the bank name as caption.")
    col1, col2, col3 = st.columns([2, 1, 1])
    with col1:
        chat = st.text_input("WhatsApp contact or group name (exactly as it appears in WhatsApp)", key="kyc_mailer_wa_chat")
    with col2:
        officer = st.text_input(
            "Sender (WhatsApp login)",
            key="kyc_mailer_wa_officer",
            help="Each officer keeps their own WhatsApp login. Type the same name every time, e.g. your first name.",
        )
    with col3:
        if st.button("Open WhatsApp & log in", key="kyc_mailer_wa_login"):
            status = st.empty()
            with st.spinner("Opening WhatsApp Web..."):
                code = run_whatsapp(["login"], lambda event: status.info(event.get("message", event.get("status", ""))), officer)
            status.success("WhatsApp is logged in on this PC.") if code == 0 else status.error(
                "WhatsApp login was not completed. Scan the QR code in the window and try again."
            )

    if st.button("Show my WhatsApp chat names", key="kyc_mailer_wa_chats"):
        names: List[str] = []
        with st.spinner("Reading the WhatsApp chat list..."):
            run_whatsapp(["chats"], lambda event: names.extend(event.get("names", [])), officer)
        st.write(names or "No chat names were read. Log in to WhatsApp first.")

    chosen = st.multiselect(
        "Screenshots to send",
        [path.name for path in existing],
        default=[path.name for path in existing],
        key="kyc_mailer_wa_files",
    )
    confirm = st.checkbox(f"Send {len(chosen)} screenshot(s) to \"{chat}\" on WhatsApp", key="kyc_mailer_wa_confirm")
    if st.button("Send on WhatsApp", type="primary", disabled=not (chat and chosen and confirm), key="kyc_mailer_wa_send"):
        items = [
            {"image": str(folder / name), "caption": name.replace(" - sent mail.png", "")}
            for name in chosen
        ]
        job_file = folder / "whatsapp_job.json"
        job_file.write_text(json.dumps({"chat": chat, "items": items}), encoding="utf-8")
        progress, log_box = st.progress(0.0), st.empty()
        events: List[Dict] = []

        def handle(event: Dict) -> None:
            events.append(event)
            done = sum(1 for item in events if item.get("status") in {"sent", "failed"})
            progress.progress(min(done / len(items), 1.0), text=f"{done}/{len(items)} {event.get('caption', '')}")
            log_box.write(event)

        code = run_whatsapp(["send", str(job_file)], handle, officer)
        failed = [event for event in events if event.get("status") == "failed"]
        blocked = [event for event in events if event.get("status") == "not_logged_in"]
        missing_chat = [event for event in events if event.get("status") == "chat_not_found"]
        if blocked:
            st.error("WhatsApp is not logged in for this sender. Use 'Open WhatsApp & log in' first.")
        elif missing_chat:
            st.error(missing_chat[0].get("error", "The WhatsApp chat was not found."))
        elif code == 0 and not failed:
            st.success(f"{len(items)} screenshot(s) sent to \"{chat}\" on WhatsApp.")
        else:
            st.error(f"{len(failed)} of {len(items)} could not be sent. Details above.")
        job_file.unlink(missing_ok=True)


def render_kyc_notice_mailer_page():
    if not _require_page_password():
        return
    st.title("KYC Notice Mailer (NIC Mail)")
    _render_password_controls()

    # Notice type selector -------------------------------------------------
    notice_type = st.radio(
        "Notice type",
        [NOTICE_TYPE_KYC, NOTICE_TYPE_PR, NOTICE_TYPE_PRABODH],
        horizontal=True,
        key="kyc_mailer_notice_type",
    )
    is_payer_remitter = notice_type == NOTICE_TYPE_PR
    is_prabodh = notice_type == NOTICE_TYPE_PRABODH

    if is_prabodh:
        st.caption("Sends the KYC Prabodh notice per bank with its own format and subject. "
                   "Upload an Excel with BANK NAME and ACCOUNT NUMBER (or AC NO). "
                   "Attachments: PDF notice + KYC Excel. The bank e-mail directory is shared.")
    elif is_payer_remitter:
        st.caption(
            "Sends one Payer/Remitter bank statement request mail per bank through "
            "the NIC mail server. The mail body is the notice without the stamp, "
            "signature and officer name. Attachments: PDF notice + Excel "
            "(the Word file is never attached)."
        )
    else:
        st.caption(
            "Sends one KYC notice mail per bank through the NIC mail server. The "
            "mail body is the notice without the stamp, signature and officer name. "
            "Attachments: PDF notice + KYC Excel (the Word file is never attached)."
        )

    # 1. Files -------------------------------------------------------------
    col1, col2 = st.columns(2)
    with col1:
        kyc_file = st.file_uploader(
            "Prabodh accounts Excel" if is_prabodh else "Pending KYC Excel (e.g. SEP PENDING KYC.xlsx)",
            type=["xlsx", "xls"], key="kyc_mailer_prabodh_upload" if is_prabodh else "kyc_mailer_kyc")
    with col2:
        email_file = st.file_uploader("Bank-wise e-mail IDs Excel", type=["xlsx", "xls"], key="kyc_mailer_emails")
        if email_file is not None:
            email_name, email_bytes = email_file.name, email_file.getvalue()
        elif DEFAULT_EMAIL_SHEET.exists():
            email_name, email_bytes = DEFAULT_EMAIL_SHEET.name, DEFAULT_EMAIL_SHEET.read_bytes()
            st.caption(f"Using {DEFAULT_EMAIL_SHEET.name} from the project folder. Upload a file to use a different one.")
        else:
            email_name, email_bytes = None, None

    col1, col2, col3 = st.columns([1, 2, 2])
    with col1:
        notice_date = st.date_input("Notice date", value=_current_notice_date(), format="DD-MM-YYYY", key="kyc_mailer_date")
    with col2:
        period = st.text_input("Subject bracket text", value="September 1 to 15", key="kyc_mailer_period",
                               help="Subject becomes: <notice subject> (<this text> - <bank name>)")

    if kyc_file is None or email_bytes is None:
        st.info("Upload both files to prepare the mails.")
        return

    with st.spinner("Preparing bank-wise notices..."):
        try:
            builder = _cached_prabodh_bundles if is_prabodh else _cached_pr_bundles if is_payer_remitter else _cached_bundles
            bundles = [b for b in builder(kyc_file.getvalue(), notice_date) if b.get("files")]
        except (ValueError, KeyError, FileNotFoundError) as exc:
            st.error(str(exc))
            return
    if not bundles:
        st.warning("No pending accounts found in this file.")
        return

    first_docx = next(data for name, data in bundles[0]["files"].items() if name.endswith(".docx"))
    default_subject = DEFAULT_PR_SUBJECT if is_payer_remitter else notice_subject_from_docx(first_docx)
    with col3:
        subject_base = st.text_input("Notice subject", value=default_subject, key=f"kyc_mailer_subject_{notice_type}")

    email_sheets = pd.ExcelFile(io.BytesIO(email_bytes)).sheet_names
    email_df = pd.read_excel(io.BytesIO(email_bytes), sheet_name=email_sheets[0], dtype=str, keep_default_na=False)
    bank_col, email_col = detect_directory_columns(list(email_df.columns))
    col1, col2 = st.columns(2)
    columns = list(email_df.columns)
    with col1:
        bank_col = st.selectbox("Bank name column", columns, index=columns.index(bank_col) if bank_col in columns else 0, key="kyc_mailer_bank_col")
    with col2:
        email_col = st.selectbox("E-mail IDs column", columns, index=columns.index(email_col) if email_col in columns else min(1, len(columns) - 1), key="kyc_mailer_email_col")
    directory = load_email_directory(email_df, bank_col, email_col)

    # 2. Review recipients -------------------------------------------------
    sent_subjects = already_sent_subjects()
    review_rows = []
    for bundle in bundles:
        match = match_bank_emails(bundle["bank_name"], directory)
        subject = build_subject(subject_base, period, bundle["bank_name"])
        already = subject in sent_subjects
        review_rows.append(
            {
                "Send": bool(match["emails"]) and not already,
                "Bank": bundle["bank_name"],
                "Accounts": bundle["accounts"],
                "Matched in e-mail sheet": match["matched_name"],
                "Match": "Already sent" if already else match["match"],
                "E-mail IDs": "; ".join(match["emails"]),
            }
        )

    st.subheader("Review recipients")
    st.caption("Check every 'Closest' match. You can edit the e-mail IDs (separate with ;) and untick banks to skip.")
    edited = st.data_editor(
        pd.DataFrame(review_rows),
        hide_index=True,
        width="stretch",
        disabled=["Bank", "Accounts", "Matched in e-mail sheet", "Match"],
        column_config={
            "Send": st.column_config.CheckboxColumn(width="small"),
            "Accounts": st.column_config.NumberColumn(width="small"),
            "E-mail IDs": st.column_config.TextColumn(width="large"),
        },
        key=f"kyc_mailer_review_{hash((notice_type, kyc_file.getvalue(), email_bytes, period, subject_base, str(notice_date)))}",
    )
    not_found = int((edited["E-mail IDs"].map(lambda v: not extract_emails(v))).sum())
    closest = int(edited["Match"].str.startswith("Closest").sum())
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Banks", len(edited))
    m2.metric("Selected to send", int(edited["Send"].sum()))
    m3.metric("Closest matches to check", closest)
    m4.metric("No e-mail found", not_found)

    bundle_by_bank = {bundle["bank_name"]: bundle for bundle in bundles}

    # 3. Preview -----------------------------------------------------------
    st.subheader("Preview")
    preview_bank = st.selectbox("Bank", list(edited["Bank"]), key="kyc_mailer_preview_bank")
    preview_bundle = bundle_by_bank[preview_bank]
    preview_docx = next(d for n, d in preview_bundle["files"].items() if n.endswith(".docx"))
    preview_html, preview_images = _cached_email(preview_docx)
    st.markdown(f"**Subject:** {escape(build_subject(subject_base, period, preview_bank))}")
    st.markdown("**Attachments:** " + ", ".join(escape(n) for n in mail_attachments(preview_bundle)))
    rendered = preview_html
    for cid, (subtype, data) in preview_images.items():
        rendered = rendered.replace(f"cid:{cid}", f"data:image/{subtype};base64,{base64.b64encode(data).decode()}")
    components.html(f'<div style="background:#fff;padding:16px">{rendered}</div>', height=650, scrolling=True)
    for name, data in mail_attachments(preview_bundle).items():
        st.download_button(f"Download {name}", data, file_name=name, key=f"kyc_mailer_dl_{name}")

    # 4. Send --------------------------------------------------------------
    st.subheader("Send")
    col1, col2, col3 = st.columns(3)
    with col1:
        sender = st.text_input("NIC e-mail ID (sender)", value=_secret("sender"), key="kyc_mailer_sender")
    with col2:
        password = st.text_input("Password / app-specific password", value=_secret("password"), type="password", key="kyc_mailer_password")
    with col3:
        cc_text = st.text_input("CC (optional)", value=_secret("cc"), key="kyc_mailer_cc")
    sender_name = st.text_input(
        "Sender name shown on the record screenshot",
        value=_secret("sender_name") or DEFAULT_SENDER_NAME,
        key="kyc_mailer_sender_name",
    )
    delay_seconds = st.slider("Pause between mails (seconds)", 0, 30, 5, key="kyc_mailer_delay")
    cc = extract_emails(cc_text)

    col1, col2 = st.columns(2)
    with col1:
        if st.button("Test login + send preview bank's mail to myself", width="stretch"):
            if not sender or not password:
                st.error("Enter the sender e-mail and password.")
            else:
                try:
                    with st.spinner("Sending test mail..."):
                        connection = open_smtp(sender, password)
                        message = build_message(
                            sender, [sender], "[TEST] " + build_subject(subject_base, period, preview_bank),
                            preview_html, preview_images, mail_attachments(preview_bundle),
                        )
                        connection.send_message(message)
                        connection.quit()
                    st.success(f"Test mail sent to {sender}. Check how it looks before sending to banks.")
                except smtplib.SMTPAuthenticationError as exc:
                    st.error(_login_failed_message(exc))
                except Exception as exc:
                    st.error(f"Test failed: {exc}")

    selected = edited[edited["Send"]]
    missing = [bank for bank, emails in zip(selected["Bank"], selected["E-mail IDs"]) if not extract_emails(emails)]
    with col2:
        confirm = st.checkbox(f"I have reviewed the recipients - send {len(selected)} mail(s) to the banks", key="kyc_mailer_confirm")
        send_clicked = st.button("Send to banks", type="primary", disabled=not confirm or selected.empty, width="stretch")

    if missing:
        st.warning("Selected banks without any e-mail ID will be skipped: " + ", ".join(missing))
    unticked = edited[~edited["Send"] & (edited["Match"] != "Already sent")]
    if not unticked.empty:
        st.warning(
            f"{len(unticked)} bank(s) are NOT ticked and will not get a mail: "
            + ", ".join(unticked["Bank"])
            + ". Tick them in the review table if this was not intended."
        )

    if send_clicked:
        if not sender or not password:
            st.error("Enter the sender e-mail and password.")
            return
        try:
            connection = open_smtp(sender, password)
        except smtplib.SMTPAuthenticationError as exc:
            st.error(_login_failed_message(exc))
            return
        except Exception as exc:
            st.error(f"Could not connect to {SMTP_HOST}: {exc}")
            return

        progress = st.progress(0.0)
        results = []
        sent_subjects = already_sent_subjects()
        for position, row in enumerate(selected.to_dict("records"), start=1):
            bank = row["Bank"]
            recipients = extract_emails(row["E-mail IDs"])
            subject = build_subject(subject_base, period, bank)
            bundle = bundle_by_bank[bank]
            attachments = mail_attachments(bundle)
            entry = {
                "sent_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "bank_name": bank,
                "subject": subject,
                "recipients": "; ".join(recipients),
                "attachments": "; ".join(attachments),
                "sender": sender,
                "sender_name": sender_name,
            }
            if not recipients:
                entry.update(status="skipped", error="No e-mail ID")
            elif subject in sent_subjects:
                entry.update(status="skipped", error="Already sent earlier")
            else:
                docx_bytes = next(d for n, d in bundle["files"].items() if n.endswith(".docx"))
                html, images = _cached_email(docx_bytes)
                message = build_message(sender, recipients, subject, html, images, attachments, cc)
                for attempt in range(2):
                    try:
                        connection.send_message(message)
                        entry.update(status="sent", error="")
                        sent_subjects.add(subject)
                        break
                    except (smtplib.SMTPServerDisconnected, smtplib.SMTPSenderRefused, ConnectionError) as exc:
                        entry.update(status="failed", error=str(exc))
                        if attempt == 0:
                            try:
                                connection = open_smtp(sender, password)
                            except Exception as reconnect_exc:
                                entry.update(error=f"Reconnect failed: {reconnect_exc}")
                                break
                    except Exception as exc:
                        entry.update(status="failed", error=str(exc))
                        break
            if entry["status"] != "skipped":
                append_send_log(entry)
            results.append(entry)
            progress.progress(position / len(selected), text=f"{position}/{len(selected)}  {bank}: {entry['status']}")
            if entry["status"] == "sent" and position < len(selected) and delay_seconds:
                time.sleep(delay_seconds)
        try:
            connection.quit()
        except Exception:
            pass

        result_df = pd.DataFrame(results)
        counts = result_df["status"].value_counts().to_dict()
        st.success(f"Sent: {counts.get('sent', 0)}   Failed: {counts.get('failed', 0)}   Skipped: {counts.get('skipped', 0)}")
        st.dataframe(result_df, hide_index=True, width="stretch")
        st.download_button("Download send report (CSV)", result_df.to_csv(index=False).encode("utf-8-sig"),
                           file_name=f"kyc_mail_report_{datetime.now():%Y%m%d_%H%M}.csv")

        sent_entries = [entry for entry in results if entry["status"] == "sent"]
        if sent_entries:
            with st.spinner(f"Saving {len(sent_entries)} sent-mail screenshot(s)..."):
                try:
                    saved = save_sent_mail_screenshots(
                        sent_entries, bundle_by_bank, subject_base, period, notice_date
                    )
                    saved_folder = records_folder(notice_date) / "KYC Prabodh" if is_prabodh else records_folder(notice_date)
                    st.success(f"{len(saved)} screenshot(s) saved in {saved_folder}")
                except Exception as exc:
                    st.error(f"Could not save the screenshots: {exc}")

    _render_records_section(bundle_by_bank, subject_base, period, notice_date, sender_name, sender, kyc_file.getvalue(), is_payer_remitter, is_prabodh)

    with st.expander("Send history"):
        st.dataframe(read_send_log(), hide_index=True, width="stretch")
