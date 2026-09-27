"""Prabodh notices, kept separate from the existing KYC notice builders."""

from copy import deepcopy
from datetime import date
import io
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from src.bankwise_kyc_notices import bank_excel_rows, bank_folder_name
from src.gujarat_account_formatter import (
    SOURCE_FIELDS, _normalize_header, _set_cant_split, _set_notice_table_cell,
    dataframe_to_styled_excel_bytes, prepare_kyc_notice_accounts,
)

TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "kyc prabodh.docx"


def prepare_prabodh_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Accept bank/account uploads without requiring pending-KYC or ACK fields."""
    df = df.copy()
    headers = {_normalize_header(column): column for column in df.columns}
    for target, aliases in SOURCE_FIELDS.items():
        if target not in df.columns:
            source = next((headers[a] for a in aliases if a in headers), None)
            if source is not None:
                df[target] = df[source]
    missing = [c for c in ("BANK NAME", "AC NO") if c not in df.columns]
    if missing:
        raise ValueError("Prabodh upload requires BANK NAME and ACCOUNT NUMBER (or AC NO).")
    for column in ("ACK NO", "IFSC CODE", "ACCOUNT HOLDER'S NAME",
                   "ACCOUNT HOLDER'S MOBILE NUMBER", "ACCOUNT HOLDER'S ADDRESS",
                   "ACCOUNT HOLDER'S LOCATION"):
        if column not in df.columns:
            df[column] = ""
    return df


def build_prabodh_docx(notice_df: pd.DataFrame, notice_date: date, bank_name: str) -> bytes:
    document = Document(TEMPLATE)
    paragraphs = document.paragraphs
    to_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "To,")
    addressee = next(p for p in paragraphs[to_index + 1:] if p.text.strip())
    addressee.runs[0].text = f"The Nodal Officer,\n{bank_name}"
    for run in addressee.runs[1:]:
        run.text = ""
    date_line = next(p for p in paragraphs if p.text.strip().startswith("Date:"))
    date_line.runs[0].text = f"Date: {notice_date:%d/%m/%Y}"
    for run in date_line.runs[1:]:
        run.text = ""
    table = document.tables[0]
    sample = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for record in notice_df.to_dict("records"):
        table._tbl.append(deepcopy(sample))
        row = table.rows[-1]
        _set_cant_split(row)
        _set_notice_table_cell(row.cells[0], record["BANK NAME"], WD_ALIGN_PARAGRAPH.LEFT)
        _set_notice_table_cell(row.cells[1], record["AC NO"], WD_ALIGN_PARAGRAPH.CENTER)
    document.core_properties.title = f"KYC Prabodh {notice_date:%d%m%Y} - {bank_name}"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def prabodh_docx_to_pdf(docx_bytes: bytes) -> bytes:
    """Render the populated template text, images and two-column account table."""
    document = Document(io.BytesIO(docx_bytes))
    output = io.BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=A4, leftMargin=45, rightMargin=40,
                           topMargin=20, bottomMargin=30)
    styles = {
        align: ParagraphStyle(f"Prabodh{align}", fontName="Times-Roman", fontSize=12,
                              leading=15, alignment=align, spaceAfter=7)
        for align in (TA_LEFT, TA_CENTER, TA_RIGHT)
    }
    story = []
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            jc = element.find(f"{qn('w:pPr')}/{qn('w:jc')}")
            align = {"center": TA_CENTER, "right": TA_RIGHT}.get(
                jc.get(qn("w:val")) if jc is not None else "left", TA_LEFT)
            for drawing in element.iter(qn("w:drawing")):
                blip = next(drawing.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"), None)
                extent = next(drawing.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"), None)
                if blip is not None:
                    part = document.part.related_parts[blip.get(qn("r:embed"))]
                    width = int(extent.get("cx")) / 12700 if extent is not None else 70
                    height = int(extent.get("cy")) / 12700 if extent is not None else 70
                    scale = min(1, pdf.width / width)
                    picture = Image(io.BytesIO(part.blob), width=width * scale, height=height * scale)
                    picture.hAlign = {TA_LEFT: "LEFT", TA_CENTER: "CENTER", TA_RIGHT: "RIGHT"}[align]
                    story.append(picture)
            text = []
            for run in element.iter(qn("w:r")):
                value = "".join(escape(child.text or "") if child.tag == qn("w:t")
                                else "<br/>" if child.tag in (qn("w:br"), qn("w:cr"))
                                else " " if child.tag == qn("w:tab") else ""
                                for child in run.iterchildren())
                props = run.find(qn("w:rPr"))
                if props is not None and props.find(qn("w:b")) is not None:
                    value = f"<b>{value}</b>"
                text.append(value)
            if "".join(text).strip():
                story.append(Paragraph("".join(text), styles[align]))
        elif element.tag == qn("w:tbl"):
            rows = [[Paragraph(escape("".join(t.text or "" for t in cell.iter(qn("w:t")))),
                               styles[TA_LEFT]) for cell in row.findall(qn("w:tc"))]
                    for row in element.findall(qn("w:tr"))]
            table = LongTable(rows, colWidths=[pdf.width * .5] * 2, repeatRows=1)
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.grey),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                       ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            story.extend([table, Spacer(1, 10)])
    pdf.build(story)
    return output.getvalue()


def build_prabodh_bundles(df: pd.DataFrame, notice_date: date) -> list[dict]:
    rows = prepare_prabodh_rows(df)
    bundles = []
    for bank_name, bank_rows in rows.groupby("BANK NAME", sort=True):
        notice_df, stats = prepare_kyc_notice_accounts(bank_rows)
        if notice_df.empty:
            continue
        name = bank_folder_name(str(bank_name))
        docx = build_prabodh_docx(notice_df, notice_date, str(bank_name))
        bundles.append({
            "bank_name": str(bank_name), "folder_name": name, "accounts": len(notice_df),
            "skipped_rows": stats["invalid_account_rows"] + stats["duplicate_account_rows"],
            "files": {
                f"{name} KYC Prabodh Notice.docx": docx,
                f"{name} KYC Prabodh Notice.pdf": prabodh_docx_to_pdf(docx),
                f"{name} KYC Details.xlsx": dataframe_to_styled_excel_bytes(
                    bank_excel_rows(bank_rows, notice_df), sheet_name="KYC Details"),
            },
        })
    return bundles
