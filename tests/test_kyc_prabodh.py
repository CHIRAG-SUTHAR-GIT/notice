from datetime import date
import io
from html import unescape
import re

from docx import Document
import pandas as pd
from pypdf import PdfReader
import pytest

from src.kyc_prabodh import build_prabodh_bundles, prepare_prabodh_rows
from src.kyc_notice_mailer import mail_attachments, notice_docx_to_email, notice_subject_from_docx


def test_prabodh_bank_account_upload_and_matching_attachments():
    source = pd.DataFrame({
        "Bank Name": ["Example Bank", "Example Bank", "Example Bank", "Other Bank"],
        "Account Number": ["001234567890", "001234567890", "invalid", "999876543210"],
        "ACCOUNT HOLDER'S NAME": ["Known"] * 4,
        "ACCOUNT HOLDER'S ADDRESS": ["Known"] * 4,
        "ACCOUNT HOLDER'S MOBILE NUMBER": ["9999999999"] * 4,
    })
    bundles = build_prabodh_bundles(source, date(2026, 9, 27))
    assert [b["bank_name"] for b in bundles] == ["Example Bank", "Other Bank"]
    assert [b["accounts"] for b in bundles] == [1, 1]
    bundle = bundles[0]
    docx = next(v for k, v in bundle["files"].items() if k.endswith(".docx"))
    document = Document(io.BytesIO(docx))
    assert [c.text for c in document.tables[0].rows[1].cells] == ["Example Bank", "001234567890"]
    assert len(document.tables[0].rows) == 2
    assert "Date: 27/09/2026" in "\n".join(p.text for p in document.paragraphs)
    assert notice_subject_from_docx(docx) == "Provide Complete details of the below mention Bank Account"
    html, _ = notice_docx_to_email(docx)
    body_text = unescape(re.sub("<[^>]+>", "", html)).replace("\xa0", " ")
    assert "Myself, Superintendent of Police" in body_text
    assert "within 10 minutes" in body_text
    assert "001234567890" in html
    assert "AXIS BANK" not in html and "924010011701859" not in html
    assert "Keshvala" not in html
    attachments = mail_attachments(bundle)
    assert sorted(attachments) == ["Example Bank KYC Details.xlsx", "Example Bank KYC Prabodh Notice.pdf"]
    reader = PdfReader(io.BytesIO(attachments["Example Bank KYC Prabodh Notice.pdf"]))
    pdf_text = "\n".join(p.extract_text() for p in reader.pages)
    assert "001234567890" in pdf_text and "within 10 minutes" in pdf_text
    assert "Keshvala" in pdf_text and "AXIS BANK" not in pdf_text
    excel = pd.read_excel(io.BytesIO(attachments["Example Bank KYC Details.xlsx"]), dtype=str)
    assert excel["AC NO"].tolist() == ["001234567890"]


def test_prabodh_rejects_upload_without_bank_or_account():
    with pytest.raises(ValueError, match="BANK NAME and ACCOUNT NUMBER"):
        prepare_prabodh_rows(pd.DataFrame({"Unrelated": ["123"]}))
